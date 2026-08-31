"""Word 报告渲染（R6）。

「报告生成使用 docxtpl：模板是真正的 .docx，Jinja 变量写在文档里，
设计师可直接改版式。」

R0.1 在这里的落点是**渲染上下文只来自校验过的 result.json**：
「报告模板负责把校验过的结构化数据填进版式，绝对不许二次加工数据。」

所以 build_context 做两件事：把 {"value":…, "source":…} 拍平成可直接引用的
值，以及把叙述文本先过一遍数字越界拦截。模板里拿到的 `m.nps` 就是
result.json 里那个数，中间没有任何加工环节。
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from bench.accuracy.number_guard import check_narrative
from bench.accuracy.provenance import allowed_numbers, iter_metrics

logger = logging.getLogger(__name__)


class RenderError(Exception):
    pass


class _MetricView(dict):
    """既能当值用，也能取 .source。

    模板里写 {{ m.nps }} 得到数值，写 {{ m.nps.source.file }} 得到来源 ——
    这样「关键数字支持点击查看来源」将来接上时，模板不用改。
    """

    def __init__(self, node: dict[str, Any]) -> None:
        super().__init__(node)
        self._value = node.get("value")

    def __str__(self) -> str:
        v = self._value
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v)

    def __format__(self, spec: str) -> str:
        if spec and isinstance(self._value, (int, float)):
            return format(self._value, spec)
        return str(self)

    @property
    def value(self) -> Any:
        return self._value

    @property
    def source(self) -> dict[str, Any]:
        return self.get("source", {})


def _flatten(node: Any) -> Any:
    if isinstance(node, dict):
        if "value" in node:
            return _MetricView(node)
        return {k: _flatten(v) for k, v in node.items()}
    if isinstance(node, list):
        return [_flatten(v) for v in node]
    return node


def build_context(result: dict[str, Any], *,
                  narrative: dict[str, str] | None = None,
                  extra: dict[str, Any] | None = None,
                  enforce_number_guard: bool = True) -> dict[str, Any]:
    """构造模板上下文。

    result   —— 已通过 schema 与 provenance 校验的 result.json
    narrative —— LLM 生成的叙述段落，进模板前必须过数字关
    """
    metrics = result.get("metrics") or {}
    ctx: dict[str, Any] = {
        "m": _flatten(metrics),
        "metrics": _flatten(metrics),
        "meta": result.get("meta", {}),
    }

    if narrative:
        allowed = allowed_numbers(result)
        checked: dict[str, str] = {}
        for key, text in narrative.items():
            if enforce_number_guard:
                r = check_narrative(text, allowed)
                if not r.ok:
                    # 宁可不出报告，也不出一份数字被模型改过的报告
                    raise RenderError(
                        f"叙述段落 {key!r} 未通过数字校验 — {r.message()}"
                    )
            checked[key] = text
        ctx["narrative"] = checked
        ctx["n"] = checked

    if extra:
        ctx.update(extra)
    return ctx


def render_docx(template: Path, output: Path, context: dict[str, Any],
                *, images: dict[str, Path] | None = None) -> Path:
    """渲染 .docx。images 里的 PNG 以 InlineImage 插入。"""
    try:
        from docx.shared import Mm
        from docxtpl import DocxTemplate, InlineImage
    except ImportError as exc:  # pragma: no cover
        raise RenderError("未安装 docxtpl") from exc

    if not template.exists():
        raise RenderError(f"报告模板不存在：{template}")

    doc = DocxTemplate(str(template))
    ctx = dict(context)
    for name, img in (images or {}).items():
        if img.exists():
            ctx[name] = InlineImage(doc, str(img), width=Mm(160))
        else:
            logger.warning("图表文件缺失，跳过：%s", img)
            ctx[name] = ""

    try:
        doc.render(ctx)
    except Exception as exc:  # noqa: BLE001
        raise RenderError(f"模板渲染失败：{exc}") from exc

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return output


def audit_report_numbers(result: dict[str, Any], rendered_text: str) -> list[str]:
    """成品复核：报告成文里的数字是否都能追溯到 result.json。

    成功指标「报告数字可追溯率 100%」量的就是这个。渲染前的守卫管的是
    LLM 生成的段落，这个函数管的是**最终成文**——模板里如果有人手写了
    一个硬编码数字，只有这一步能发现。
    """
    r = check_narrative(rendered_text, allowed_numbers(result))
    return r.violations


def extract_text(docx_path: Path) -> str:
    """抽出 .docx 的全部文字（含表格），供上面的复核用。"""
    from docx import Document

    doc = Document(str(docx_path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)
